"""Generates the monthly Corte de Caja workbook - an automatically
regenerated replacement for the accountant's manually-built spreadsheet,
written to a shared folder so she can keep opening it the same way she
always has instead of using the web dashboard directly (requested
2026-09-24, see the corte-de-caja-dashboard project memory).

Layout follows the real "Corte de caja 2026.xlsx" day sheets (reworked
2026-09-25 to match them closely, keeping everything the automatic version
adds): Arial 8, thin borders, merged title/long-format date, the same 11
columns (CUENTA ... OBSERVACIONES, with OBSERVACIONES merged across K:O),
then our two extra columns (FORMA DE PAGO, BANCO). Rows are grouped in
blocks with a bold header row and a bold =SUM subtotal, but by ZONE (named
from the ERP's own agent names, optionally overridden with a salesperson's
name via CORTE_DE_CAJA_ZONE_LABELS) because who physically collected the
cash isn't derivable. Like the original, transfers sit in a trailing
section outside TOTAL CORTE, and the bottom block reads TOTAL EFECTIVO
(= corte - terminal - cheques), TOTAL TERMINAL, TOTAL CHEQUES, TOTAL CORTE
(= sum of the block subtotals) - here computed with live formulas. Added on
top: TOTAL TRANSFERENCIAS, the day's grand total, per-zone totals across all
payment methods, and a section for rows excluded in the dashboard.
OBSERVACIONES is filled the way the accountant writes it (TRANSFERENCIA
BBVA, TERMINAL, the cheque note; blank means cash), with "(SUG.)" on
payment methods that are only the automatic suggestion.

Like the web dashboard, this only ever reads from the ERP + this app's own
CorteDeCajaAdjustment table - the old spreadsheet is never read as an input.
"""

import calendar
from datetime import date
from pathlib import Path

from django.conf import settings
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from api.models import AdmAgentes, AdmConceptos, AdmDocumentos
from corte_de_caja.models import CorteDeCajaAdjustment
from corte_de_caja.services import calculate_corte_de_caja

# No workbook is ever generated for a month before this one - the user was
# explicit (2026-09-24) that this replaces the manual process going
# forward, not a backfill of it.
EARLIEST_EXPORT_MONTH = date(2026, 9, 1)

HEADERS = [
    'CUENTA', 'NOMBRE DEL CLIENTE', 'No. FACTURA', 'DEBE', 'HABER', 'ENTREGA', 'RECIBE',
    'VENCIMIENTO', 'COMISION', 'EQ-R-S', 'OBSERVACIONES',
]
COL_FORMA_PAGO, COL_BANCO = 16, 17  # after OBSERVACIONES (K:O merged)
LAST_COL = COL_BANCO
HEADER_ROW = 5

ZONE_SHEET_CODES = {
    'ZONA1': 'Z-1', 'ZONA2': 'Z-2', 'OFICINA': 'O', 'SERVICIOS': 'S', 'PUNTOVENTA': 'PV',
}
ZONE_ORDER = ['ZONA1', 'ZONA2', 'OFICINA', 'SERVICIOS', 'PUNTOVENTA']
CATEGORY_SHEET_CODES = {
    'R': 'R', 'R_NW': 'R', 'R_CHEM': 'R', 'R_FAN': 'R', 'B': 'B', 'S': 'S',
}
PAYMENT_METHOD_LABELS = dict(CorteDeCajaAdjustment.PAYMENT_METHOD_CHOICES)
# How the accountant abbreviates the receiving bank in OBSERVACIONES.
BANK_ABBREVIATIONS = (('BBVA', 'BBVA'), ('HSBC', 'HSBC'), ('BANAMEX', 'BMX'), ('BANORTE', 'BTE'))

# Spanish month abbreviations, matching the real sheet's own tab naming
# (e.g. "22.Sep", "15.Ago") - built explicitly rather than via
# date.strftime('%b'), which depends on the server's system locale and
# would silently render English abbreviations (Jan/Aug/Dec) on a server
# that isn't configured with an es_MX locale.
SPANISH_MONTH_ABBR = [
    'Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic',
]

# Formats copied from the real sheet, with a currency symbol added to money.
MONEY_FORMAT = '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@_-'
DUE_DATE_FORMAT = 'd\\-mmm\\-yy'
LONG_DATE_FORMAT = '[$-F800]dddd", "mmmm\\ dd", "yyyy'

FONT = Font(name='Arial', size=8)
FONT_BOLD = Font(name='Arial', size=8, bold=True)
FONT_TITLE = Font(name='Arial', size=10, bold=True)
FONT_EXCLUDED = Font(name='Arial', size=8, color='808080', strike=True)
THIN = Side(style='thin')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
COLUMN_WIDTHS = {
    1: 12, 2: 33, 3: 10.7, 4: 7, 5: 15, 6: 7.3, 7: 6.3, 8: 11.4, 9: 8.7, 10: 6.1,
    11: 9, 12: 9, 13: 9, 14: 9, 15: 9, COL_FORMA_PAGO: 24, COL_BANCO: 26,
}


def _month_bounds(year, month):
    last_day = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last_day)


def _folio_prefixes(invoice_ids):
    """No. FACTURA on the real sheet is the tax series ('A', 'B', or 'F' for
    the default 16% series) plus the bare folio number. The series lives in
    AdmConceptos.CSERIEPOROMISION - CPREFIJOCONCEPTO is 'F' for every
    concept, so using it renders every B/A invoice as 'F' (verified against
    the accountant's Sep 2026 sheet: 272/273 rows match this way).
    """
    concepto_by_invoice = dict(
        AdmDocumentos.objects.filter(CIDDOCUMENTO__in=invoice_ids).values_list(
            'CIDDOCUMENTO', 'CIDCONCEPTODOCUMENTO',
        )
    )
    serie_by_concepto = dict(
        AdmConceptos.objects.values_list('CIDCONCEPTODOCUMENTO', 'CSERIEPOROMISION')
    )

    def display(concepto_id):
        serie = (serie_by_concepto.get(concepto_id) or '').strip().upper()
        return serie if serie in ('A', 'B') else 'F'

    return {
        invoice_id: display(concepto_id)
        for invoice_id, concepto_id in concepto_by_invoice.items()
    }


def _zone_labels():
    """Block header per zone: the ERP's own agent name (ZONA 1, OFICINA...),
    unless CORTE_DE_CAJA_ZONE_LABELS overrides it (e.g. with the salesperson
    who runs that route)."""
    labels = dict(AdmAgentes.objects.values_list('CCODIGOAGENTE', 'CNOMBREAGENTE'))
    labels.update(settings.CORTE_DE_CAJA_ZONE_LABELS)
    return labels


def _bank_abbreviation(bank):
    name = (bank or '').upper()
    for needle, abbreviation in BANK_ABBREVIATIONS:
        if needle in name:
            return abbreviation
    return None


def _observaciones(row):
    """OBSERVACIONES the way the accountant writes it: transfers name the
    bank, terminal says TERMINAL, cheques carry their note, and cash is left
    blank. A payment method that is only the automatic suggestion gets
    '(SUG.)' so it isn't read as confirmed."""
    method = row['payment_method']
    note = row['note'] or ''
    if method == CorteDeCajaAdjustment.PAYMENT_METHOD_TRANSFERENCIA:
        text = ' '.join(filter(None, ['TRANSFERENCIA', _bank_abbreviation(row['bank'])]))
    elif method == CorteDeCajaAdjustment.PAYMENT_METHOD_TERMINAL:
        text = 'TERMINAL'
    elif method == CorteDeCajaAdjustment.PAYMENT_METHOD_CHEQUE:
        text = note or 'CHEQUE'
        note = ''
    else:
        text = ''
    if text and not row['payment_method_confirmed']:
        text += ' (SUG.)'
    return ' - '.join(filter(None, [text, note])) or None


class _DaySheet:
    """Writes one day sheet, tracking the current row."""

    def __init__(self, ws):
        self.ws = ws
        self.r = HEADER_ROW + 1

    def cell(self, col, value=None, *, bold=False, fmt=None, align='left', font=None, border=True):
        cell = self.ws.cell(row=self.r, column=col, value=value)
        cell.font = font or (FONT_BOLD if bold else FONT)
        cell.alignment = Alignment(horizontal=align, vertical='center')
        if fmt:
            cell.number_format = fmt
        if border:
            cell.border = BORDER
        return cell

    def blank_table_row(self):
        for col in range(1, LAST_COL + 1):
            self.cell(col)
        self.ws.merge_cells(start_row=self.r, start_column=11, end_row=self.r, end_column=15)
        self.ws.row_dimensions[self.r].height = 13.5

    def label_row(self, text):
        self.blank_table_row()
        self.cell(2, text, bold=True, align='center')
        self.r += 1

    def subtotal_row(self, first, last):
        self.blank_table_row()
        self.cell(5, f'=SUM(E{first}:E{last})', bold=True, fmt=MONEY_FORMAT, align='right')
        row = self.r
        self.r += 1
        return row

    def payment_row(self, row, folio_prefixes, excluded=False):
        self.blank_table_row()
        font = FONT_EXCLUDED if excluded else None
        folio = f"{folio_prefixes.get(row['invoice_id'], 'F')} {int(row['folio'])}"
        self.cell(1, row['cuenta'], font=font)
        self.cell(2, row['cliente'], font=font)
        self.cell(3, folio, font=font)
        self.cell(4, 'A' if row['abono'] else None, font=font, align='center')
        self.cell(5, float(row['amount']), fmt=MONEY_FORMAT, align='right', font=font)
        self.cell(8, row['due_date'].date() if row['due_date'] else None, fmt=DUE_DATE_FORMAT, align='center', font=font)
        self.cell(9, ZONE_SHEET_CODES.get(row['zone'], row['zone']), align='center', font=font)
        self.cell(10, CATEGORY_SHEET_CODES.get(row['category'], row['category']), align='center', font=font)
        self.cell(11, _observaciones(row), font=font)
        method_label = PAYMENT_METHOD_LABELS.get(row['payment_method'], '')
        if row['payment_method'] and not row['payment_method_confirmed']:
            method_label += ' (sugerido)'
        self.cell(COL_FORMA_PAGO, method_label or None, font=font)
        self.cell(COL_BANCO, row['bank'] or None, font=font)
        self.r += 1


def _zone_sort_key(zone):
    return (ZONE_ORDER.index(zone) if zone in ZONE_ORDER else len(ZONE_ORDER), zone or '')


def _write_day_sheet(wb, day, rows, folio_prefixes, zone_labels):
    ws = wb.create_sheet(f'{day.day:02d}.{SPANISH_MONTH_ABBR[day.month - 1]}'[:31])
    sheet = _DaySheet(ws)

    # Title, long-format date, source note - same cells as the original.
    ws.merge_cells('A1:Q1')
    ws['A1'] = 'CORTE DE CAJA COBRANZA GENERAL'
    ws['A1'].font = FONT_TITLE
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    for col in range(1, LAST_COL + 1):
        ws.cell(row=1, column=col).border = BORDER
    ws.merge_cells('C2:F2')
    ws['C2'] = day
    ws['C2'].number_format = LONG_DATE_FORMAT
    ws['C2'].font = FONT_BOLD
    ws['C2'].alignment = Alignment(horizontal='center', vertical='center')
    for col in range(3, 7):
        ws.cell(row=2, column=col).border = Border(bottom=THIN)
    ws['A3'] = 'CAPTURA CONTPAQ COMERCIAL'
    ws['A3'].font = FONT_BOLD
    for r in (1, 2, 3, 4):
        ws.row_dimensions[r].height = 12.75

    sheet.r = HEADER_ROW
    sheet.blank_table_row()
    for col, title in enumerate(HEADERS, start=1):
        sheet.cell(col, title, bold=True, align='center')
    sheet.cell(COL_FORMA_PAGO, 'FORMA DE PAGO', bold=True, align='center')
    sheet.cell(COL_BANCO, 'BANCO', bold=True, align='center')
    sheet.r += 1

    def order(row):
        return (_zone_sort_key(row['zone']), row['cliente'], row['folio'])

    counted = [r for r in rows if not r['excluded']]
    excluded = [r for r in rows if r['excluded']]
    physical = [r for r in counted if r['payment_method'] != CorteDeCajaAdjustment.PAYMENT_METHOD_TRANSFERENCIA]
    transfers = [r for r in counted if r['payment_method'] == CorteDeCajaAdjustment.PAYMENT_METHOD_TRANSFERENCIA]

    # Efectivo / terminal / cheque, in one block per zone (their per-salesperson blocks).
    corte_subtotals = []
    for zone in sorted({r['zone'] for r in physical}, key=_zone_sort_key):
        sheet.label_row(zone_labels.get(zone, zone))
        first = sheet.r
        for row in sorted((r for r in physical if r['zone'] == zone), key=order):
            sheet.payment_row(row, folio_prefixes)
        corte_subtotals.append(sheet.subtotal_row(first, sheet.r - 1))

    # Transfers: trailing section, outside TOTAL CORTE - same as the original.
    transfer_subtotal = None
    if transfers:
        sheet.label_row('TRANSFERENCIAS')
        first = sheet.r
        for row in sorted(transfers, key=order):
            sheet.payment_row(row, folio_prefixes)
        transfer_subtotal = sheet.subtotal_row(first, sheet.r - 1)
    data_end = sheet.r - 1

    if excluded:
        sheet.label_row('EXCLUIDOS DEL CORTE (no se suman)')
        for row in sorted(excluded, key=order):
            sheet.payment_row(row, folio_prefixes, excluded=True)

    # Totals block - labels in B, values in D:E, like the original.
    def total_line(label, formula):
        sheet.r += 0
        row = sheet.r
        ws.cell(row=row, column=2, value=label).font = FONT_BOLD
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        value = ws.cell(row=row, column=4, value=formula)
        value.font = FONT_BOLD
        value.number_format = MONEY_FORMAT
        value.alignment = Alignment(horizontal='right')
        ws.row_dimensions[row].height = 12.75
        sheet.r += 1
        return row

    forma = get_column_letter(COL_FORMA_PAGO)
    rng = lambda col: f'${col}${HEADER_ROW + 1}:${col}${data_end}'
    sheet.r += 1
    r_efectivo = sheet.r
    r_terminal, r_cheques, r_corte = r_efectivo + 1, r_efectivo + 2, r_efectivo + 3
    total_line('TOTAL EFECTIVO', f'=D{r_corte}-D{r_terminal}-D{r_cheques}')
    total_line('TOTAL TERMINAL', f'=SUMIF({rng(forma)},"Terminal*",{rng("E")})')
    total_line('TOTAL CHEQUES', f'=SUMIF({rng(forma)},"Cheque*",{rng("E")})')
    total_line('TOTAL CORTE', '=SUM(' + ','.join(f'E{r}' for r in corte_subtotals) + ')' if corte_subtotals else 0)
    r_transf = total_line('TOTAL TRANSFERENCIAS', f'=E{transfer_subtotal}' if transfer_subtotal else 0)
    total_line('TOTAL COBRADO DEL DIA', f'=D{r_corte}+D{r_transf}')
    unclassified = [r for r in physical if not r['payment_method']]
    if unclassified:
        total_line('(incluye sin forma de pago, contado como efectivo)',
                   f'=SUMPRODUCT(({rng("C")}<>"")*({rng(forma)}="")*{rng("E")})')

    sheet.r += 1
    ws.cell(row=sheet.r, column=2, value='TOTALES POR ZONA (todas las formas de pago)').font = FONT_BOLD
    sheet.r += 1
    zone_col = get_column_letter(9)
    for zone in sorted({r['zone'] for r in counted}, key=_zone_sort_key):
        code = ZONE_SHEET_CODES.get(zone, zone)
        total_line(f'TOTAL {code}', f'=SUMIF({rng(zone_col)},"{code}",{rng("E")})')

    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f'A{HEADER_ROW + 1}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 1
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f'{HEADER_ROW}:{HEADER_ROW}'


def build_month_workbook(year, month):
    month_start = date(year, month, 1)
    if month_start < EARLIEST_EXPORT_MONTH:
        earliest_label = SPANISH_MONTH_ABBR[EARLIEST_EXPORT_MONTH.month - 1]
        raise ValueError(
            f'No se generan reportes anteriores a {earliest_label} {EARLIEST_EXPORT_MONTH.year} - '
            'este reemplazo automático inicia desde ese mes, no reconstruye el histórico.'
        )

    date_from, date_to = _month_bounds(year, month)
    result = calculate_corte_de_caja(date_from, date_to)

    rows_by_day = {}
    for row in result['rows']:
        rows_by_day.setdefault(row['event_date'], []).append(row)

    invoice_ids = {row['invoice_id'] for row in result['rows']}
    folio_prefixes = _folio_prefixes(invoice_ids)
    zone_labels = _zone_labels()

    wb = Workbook()
    wb.remove(wb.active)  # openpyxl creates a default blank sheet - unused, this workbook is all named day sheets
    for day in sorted(rows_by_day):
        _write_day_sheet(wb, day, rows_by_day[day], folio_prefixes, zone_labels)

    if not wb.sheetnames:
        # No payments at all this month (e.g. exporting the current month
        # before any data exists yet) - still produce a valid, openable
        # workbook rather than erroring, so the scheduled job never fails
        # just because a month is quiet so far.
        wb.create_sheet('Sin datos')

    return wb


def export_month(year, month, output_dir=None):
    """Writes (overwrites) the month's workbook to CORTE_DE_CAJA_EXPORT_DIR
    (or `output_dir`) - see core/settings.py. That directory is expected to
    be an SMB-shared folder on the office server so the accountant/manager
    can open the file directly, without any manual export step; making the
    directory actually SMB-visible is a one-time server/IT configuration,
    not something this function does.
    """
    output_dir = Path(output_dir or settings.CORTE_DE_CAJA_EXPORT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)

    wb = build_month_workbook(year, month)
    filename = f'Corte de Caja {year}-{month:02d}.xlsx'
    path = output_dir / filename

    # Write to a temp file then replace atomically, so a job that runs
    # while someone has the file open (or another export overlaps) never
    # leaves a half-written, corrupt workbook in the shared folder. If she
    # has the file open in Excel when this runs, Windows may refuse the
    # replace outright (a held file lock, not just a race) - that raises
    # here rather than corrupting anything; the temp file is simply
    # overwritten and retried on the next scheduled run.
    tmp_path = output_dir / f'.{filename}.tmp'
    wb.save(tmp_path)
    tmp_path.replace(path)
    return path
