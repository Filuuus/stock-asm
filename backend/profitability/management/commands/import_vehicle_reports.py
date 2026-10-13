from datetime import date
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from profitability.models import VehicleMonthlyCost
from profitability.vehicle_reports import parse_fuel, parse_maintenance


class Command(BaseCommand):
    help = (
        "Imports fuel and maintenance per vehicle per month from the accountant's workbooks "
        '(VEHICLE_REPORTS_DIR/<year>/Reportes Mensuales <year>.xlsx). Only finished months; re-running '
        'replaces those months. Defaults to this year and last year.'
    )

    def add_arguments(self, parser):
        parser.add_argument('--year', type=int, action='append', help='Repeat for several years')

    def handle(self, *args, **options):
        if not settings.VEHICLE_REPORTS_DIR:
            raise CommandError('Falta VEHICLE_REPORTS_DIR en backend/.env (carpeta de Contabilidad).')
        today = date.today()
        before = f'{today:%Y-%m}'
        for year in options['year'] or [today.year - 1, today.year]:
            path = Path(settings.VEHICLE_REPORTS_DIR) / str(year) / f'Reportes Mensuales {year}.xlsx'
            if not path.exists():
                self.stdout.write(f'{year}: no existe {path}')
                continue
            book = openpyxl.load_workbook(path, data_only=True, read_only=True)
            fuel = parse_fuel(book['Combustible'].iter_rows(values_only=True), year, before)
            months = {m for m, _ in fuel}
            maintenance = {k: v for k, v in parse_maintenance(
                book['Mantenimiento de Vehículos'].iter_rows(values_only=True), year, before).items() if k[0] in months}
            with transaction.atomic():
                VehicleMonthlyCost.objects.filter(month__in=months).delete()
                VehicleMonthlyCost.objects.bulk_create(
                    [VehicleMonthlyCost(month=m, vehicle=v, kind='fuel', amount=a) for (m, v), a in fuel.items()]
                    + [VehicleMonthlyCost(month=m, vehicle=v, kind='maintenance', amount=a)
                       for (m, v), a in maintenance.items()]
                )
            self.stdout.write(f'{year}: {len(months)} meses, {len(fuel)} cargas de combustible, '
                              f'{len(maintenance)} de mantenimiento.')
